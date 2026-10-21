import json
from itertools import islice
from pathlib import Path
import openpyxl
from docx import Document

FOLDER = Path('C:/Users/123/Documents/xwechat_files/wxid_1980ghh2hh3u22_dbe8/msg/file/2026-10')
files = ['规则库和专家补充清单(1).xlsx','不重要的数据_默认缺省 .xlsx','核心数据表（人工复核+ai模拟）.xlsx']
for filename in files:
    book = openpyxl.load_workbook(FOLDER/filename,read_only=True,data_only=False)
    print(filename)
    for sheet in book:
        print(json.dumps({'sheet':sheet.title,'rows':sheet.max_row,'cols':sheet.max_column,
            'first_rows':[[str(value)[:220] if value is not None else None for value in row]
                for row in islice(sheet.iter_rows(values_only=True),10)]},ensure_ascii=False))
    book.close()
template = Path('C:/Users/123/Documents/xwechat_files/wxid_1980ghh2hh3u22_dbe8/msg/attach/ccea273efb53072f324cfc379dc597b1/2026-10/Rec/7f2421bbe6e83b5a/F/3/说明文档模版.docx')
document = Document(template)
print('DOCX TEMPLATE')
print('\n'.join(p.text for p in document.paragraphs if p.text))
for i,table in enumerate(document.tables,1):
    print(json.dumps({'table':i,'rows':[[c.text for c in r.cells] for r in table.rows]},ensure_ascii=False))
