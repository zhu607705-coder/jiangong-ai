from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path('D:/建工ai')
TARGET = ROOT / 'docs/老旧小区改造案例拆分与六维数据说明.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.66)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)
section.footer_distance = Inches(.3)


def set_font(style, size, east_asia='宋体', bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), east_asia)


set_font(doc.styles['Normal'], 11.5)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.23
normal.space_after = Pt(6)
normal.widow_control = True
set_font(doc.styles['Title'], 18, '微软雅黑', True)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
set_font(doc.styles['Subtitle'], 11.5, '微软雅黑')
doc.styles['Subtitle'].paragraph_format.space_after = Pt(10)
set_font(doc.styles['Heading 1'], 13, '微软雅黑', True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
set_font(doc.styles['Heading 2'], 11.5, '微软雅黑', True)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
for style_name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[style_name]
    style.paragraph_format.keep_with_next = True
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in list(ppr.findall(qn('w:pBdr'))):
            ppr.remove(border)


def body(text):
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(23)
    return paragraph


def hyperlink(paragraph, label, url):
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), relationship)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '245070')
    props.append(color)
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Calibri')
    fonts.set(qn('w:hAnsi'), 'Calibri')
    fonts.set(qn('w:eastAsia'), '宋体')
    props.append(fonts)
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '19')
    props.append(size)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


doc.add_paragraph('老旧小区改造案例拆分与六维数据说明', 'Title')
doc.add_paragraph('建工 AI 研究说明  以求是村已有材料为例', 'Subtitle')
body('我们围绕老旧小区改造，将已有项目中的具体举措整理为可检索、可比较的经验卡，再按六个维度组织相关数据。希望请老师审阅案例拆分的尺度和六维数据的设置，帮助确定影响改造措施选择的关键因素与专业条件。')

doc.add_heading('案例如何拆分并沉淀', level=1)
body('以单项改造举措为最小整理单元。先识别“在什么对象上采取了什么动作”，再关联背景、实施条件、投入、服务对象、实施过程、运行结果和维护安排。电梯加装、供水设施更新、坡道设置和物业托管分别形成经验条目；同一举措的多段文字、表格和附件通过原文位置关联。')
body('每张经验卡保存项目及楼栋单元、举措类型、面积与数量、涉及住户、费用类型和单位、参与主体、阶段、结果、来源位置及版本。整体工程的范围与金额保存在工程层，分项投入按原文对应到具体措施。新资料补充后更新相关字段，并保留历次记录。')

doc.add_heading('求是村二次供水改造', level=2)
body('2023年8月资料记录，拟改范围为29、30、31、48幢，共252户，涉及水泵设备更新、管道材质提升和水表改造。整理时，以该工程建立父级记录，再按具体动作关联措施条目，保存楼栋、工程范围、现状问题和技术条件。阶段记录为方案公示与进场准备；留博楼的泵房选址问题单列为技术子题。[1]')

doc.add_heading('求是村电梯托管', level=2)
body('2023年12月资料记录，多方协商形成物业托管机制，并商讨托管协议、费用和管理细则。该举措保存为运行管理经验，关联社区、业委会、物业及居民代表的职责。资金字段记录当时托管费由600元降至300元，单位为元/台/月；维护字段关联检查、保养和运行记录。[2]')

doc.add_heading('求是社区邻里中心', level=2)
body('2025年10月资料记录邻里中心揭牌启用，包含康养、文化活动、儿童与家庭服务。按服务功能拆分时，各条目关联设施空间、服务对象、服务内容及承接主体。后续调查继续补充开放时间、服务容量、使用频率和居民反馈，形成从功能配置到运行表现的连续记录。[3]')

doc.add_page_break()
doc.add_heading('六维分别沉淀哪些数据', level=1)
body('六维数据围绕具体举措建立关联。同一举措可对应多个维度，例如电梯托管同时涉及实施协同、长效维护和资金交付。面积、人数、数量与费用均保存对象范围、单位、统计时点和来源。')

records = [
    ('居民需求', '户数与人数、年龄和行动能力、具体生活困难、需求调查、居民意见与分歧。', '谁需要改善\n哪些需求优先'),
    ('公共服务', '设施类型与位置、服务对象、容量、开放时间、费用、步行可达性与覆盖情况。', '现有服务怎样使用\n哪些群体需要补覆盖'),
    ('居住工程', '建筑年代、楼层与面积、设施现状、检测问题、改造位置、工程量与技术条件。', '具体改哪里\n工程方案怎样选择'),
    ('实施协同', '产权与实施范围、参与主体及职责、居民协商、公示、施工组织与分期安排。', '谁组织实施\n怎样协调各方'),
    ('长效维护', '接收与维护主体、维护内容、合同、运行费用、故障记录与交接机制。', '建成后谁负责\n怎样持续运行'),
    ('资金交付', '分项预算、资金来源、居民分摊、落实状态、支付节点、验收与结算。', '各项投入多少\n资金与交付怎样衔接'),
]
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(.9), Inches(4.4), Inches(1.6)]
for column, width in zip(table.columns, widths):
    column.width = width
for i, title in enumerate(['维度', '主要沉淀的数据', '用于回答的问题']):
    table.rows[0].cells[i].text = title
for row in records:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text

properties = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
    border = OxmlElement('w:' + edge)
    border.set(qn('w:val'), 'single')
    border.set(qn('w:sz'), '4')
    border.set(qn('w:color'), 'CBD3D9')
    borders.append(border)
properties.append(borders)
for row_index, row in enumerate(table.rows):
    row_props = row._tr.get_or_add_trPr()
    no_split = OxmlElement('w:cantSplit')
    row_props.append(no_split)
    if row_index == 0:
        repeat = OxmlElement('w:tblHeader')
        row_props.append(repeat)
    for column_index, cell in enumerate(row.cells):
        cell.width = widths[column_index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_props = cell._tc.get_or_add_tcPr()
        margins = OxmlElement('w:tcMar')
        for side, amount in [('top', 85), ('bottom', 85), ('left', 100), ('right', 100)]:
            child = OxmlElement('w:' + side)
            child.set(qn('w:w'), str(amount))
            child.set(qn('w:type'), 'dxa')
            margins.append(child)
        cell_props.append(margins)
        if row_index == 0 or row_index % 2 == 0:
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DFEAF1' if row_index == 0 else 'F4F7F9')
            cell_props.append(shade)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 or row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.16
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.bold = row_index == 0 or column_index == 0

doc.add_heading('请老师复核的重点', level=1)
body('希望老师帮助判断：单项举措的拆分尺度是否合适，六维是否覆盖规划研究需要的信息，哪些因素属于必须满足的专业条件，哪些因素会改变措施选择和实施顺序。随后将因素的定义、适用场景和判断条件整理成规则，并关联支持它的案例与文献。')
body('拟在区域研判中，将居民需求、工程条件、协商和维护安排对应到具体措施，展示采用的规则、支持资料及下一步调查事项。例如，供水改造重点比较泵房选址和管网条件；电梯托管重点比较管理主体、服务职责和费用安排，供老师检查判断与依据是否一致。')

doc.add_heading('案例资料来源', level=1)
references = [
    ('[1] 浙江在线 2023年8月18日 浙大求是社区推进二次供水改造', 'https://town.zjol.com.cn/czjsb/202308/t20230818_26113017.shtml'),
    ('[2] 杭州网转载浙江日报 2023年12月28日 西湖区三方协同为居民幸福加码', 'https://hznews.hangzhou.com.cn/shehui/content/2023-12/28/content_8666802.htm'),
    ('[3] 西湖报 2025年10月28日 这个康养生活综合体已焕新启用', 'https://xhb.xihunet.cn/h5/html5/2025-10/28/content_9248_18995526.htm'),
]
for label, url in references:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.1
    hyperlink(paragraph, label, url)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run = OxmlElement('w:r')
text = OxmlElement('w:t')
text.text = '1'
run.append(text)
field.append(run)
footer._p.append(field)
doc.core_properties.title = '老旧小区改造案例拆分与六维数据说明'
doc.core_properties.subject = '供城乡规划专业教师审阅的案例组织与六维数据说明'
doc.core_properties.author = '朱航成'
doc.core_properties.keywords = '老旧小区改造 求是村 案例拆分 六维数据 知识沉淀'
doc.core_properties.comments = ''
TARGET.parent.mkdir(parents=True, exist_ok=True)
doc.save(TARGET)
print(str(TARGET))
