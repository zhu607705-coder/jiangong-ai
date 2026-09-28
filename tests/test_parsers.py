from pathlib import Path
from io import BytesIO
import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from services.capabilities.parsers import parse_document


@pytest.fixture
def blobs(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    root = tmp_path / "blobs" / "PARSER-TEST"
    root.mkdir(parents=True)
    return root


def test_docx_keeps_chinese_text_and_paragraph_locator(blobs):
    file = blobs / "input.docx"
    doc = Document(); doc.add_paragraph("合成测试：本项目不整体拆除。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "建筑编号"; table.cell(0, 1).text = "租赁状态"
    table.cell(1, 0).text = "B-01"; table.cell(1, 1).text = "尚待核实"
    doc.save(file)
    result = parse_document(str(file), file.name, "DOC-1")
    assert result["status"] == "extracted"
    assert "paragraph:1" in result["text"] and "本项目不整体拆除" in result["text"]
    assert "table:1" in result["text"] and "B-01" in result["text"]


def test_xlsx_keeps_row_identity_and_does_not_execute_formulas(blobs):
    file = blobs / "input.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "合成建筑清单"
    sheet.append(["建筑编号", "面积", "备注"])
    sheet.append(["B-01", 1000, "仍在出租"])
    sheet.append(["B-02", 2000, "=1+2"])
    book.save(file)
    result = parse_document(str(file), file.name, "DOC-2")
    assert result["status"] == "extracted"
    assert "sheet:合成建筑清单/row:2" in result["text"]
    assert "建筑编号=B-01 | 面积=1000" in result["text"]
    assert "备注=3" not in result["text"]
    assert result["profile"]["tables"][0]["rows"][1]["rowNumber"] == 3


def test_pdf_keeps_page_locator_and_source_hash(blobs):
    file = blobs / "input.pdf"
    writer = canvas.Canvas(str(file)); writer.drawString(40, 700, "SYNTHETIC fixture: building B-01, no structural conclusion."); writer.showPage(); writer.drawString(40, 700, "Second evidence page."); writer.save()
    result = parse_document(str(file), file.name, "DOC-3")
    assert result["status"] == "extracted"
    assert "page:1" in result["text"] and "page:2" in result["text"]
    assert len(result["source_hash"]) == 64


def test_empty_scan_is_ocr_required_not_extracted(blobs):
    file = blobs / "scan.pdf"
    writer = canvas.Canvas(str(file)); writer.rect(50, 50, 100, 100); writer.showPage(); writer.save()
    result = parse_document(str(file), file.name, "DOC-4")
    assert result["status"] == "ocr_required" and result["text"] == ""


def test_parser_refuses_files_outside_project_blob_volume(blobs, tmp_path):
    file = tmp_path / "outside.docx"; file.write_bytes(b"test")
    with pytest.raises(ValueError, match="outside"):
        parse_document(str(file), file.name, "DOC-X")
