"""Read-only feasibility audit of the supplied workbook and existing source library.

No training, fitting, synthetic records, downloads, or source-file edits occur here.
Run with the bundled Python executable; openpyxl is the only external dependency.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_LIBRARY = ROOT / "data/public_cases/renewal-cost-library-20261013.json"
DEFAULT_OUTPUT = ROOT / "data/imports/cost-training-feasibility-20261014.json"

NUMERIC_INPUT_FIELDS = [
    "建筑面积_㎡", "户数", "楼栋数", "小区数量", "改造范围面积_㎡", "用地面积_㎡",
    "道路改造面积_㎡", "屋面施工面积_㎡", "外墙施工面积_㎡", "所披露管线长度_m",
    "道路长度_m", "明确改造建筑面积_㎡", "全小区总建筑面积_㎡", "环境整治面积_㎡",
]
TARGET_FIELDS = [
    "正式中标金额_元", "合同金额_元", "竣工结算金额_元", "控制价_元", "建安估算概算_元",
    "项目总投资或概算_元",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def serialize(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def positive(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value > 0


def province(value) -> str:
    return str(value or "unknown").removesuffix("省").removesuffix("市")


def work_type(row) -> str:
    mode = str(row.get("发包模式") or "")
    if "EPC" in mode.upper() or "设计施工" in mode or "勘察设计施工" in mode:
        return "epc"
    if "运营" in mode:
        return "construction_with_operation"
    if "施工" in mode:
        return "construction"
    return "unknown"


def scope_category(row) -> str:
    value = str(row.get("改造类型") or "")
    if value == "屋面及相关设施改造":
        return "roof_and_related_facilities"
    if value == "外墙饰面修缮":
        return "exterior_wall_finishes"
    if value == "无卫生设施提升改造":
        return "sanitary_facilities_upgrade"
    if value in ("综合整治", "综合改造"):
        return "comprehensive_renovation"
    if value == "微改造":
        return "micro_renovation"
    if value == "公共配套及综合改造":
        return "public_facilities_and_comprehensive"
    if value == "配套道路专项":
        return "road_subproject"
    return "unknown"


def project_group(row) -> str:
    # Investment code is preferred because procurement/report numbers may be lot-specific.
    # No undocumented digit stripping is used for Zhejiang/Guangdong codes.
    key = row.get("投资项目代码") or row.get("项目或报建编号") or row.get("项目及标段名称")
    return f"{province(row.get('省级地区'))}|{key}"


def year(value):
    if isinstance(value, (date, datetime)):
        return value.year
    match = re.search(r"\b(20\d{2})\b", str(value or ""))
    return int(match.group(1)) if match else None


def entry(row, excel_row: int, stage: str, amount: float) -> dict:
    area = row.get("建筑面积_㎡")
    notes = str(row.get("入模限制与核验备注") or "")
    tender_or_award_days = row.get("招标或中标工期_天")
    # Combined column is not universally a pre-award feature. The Shanghai notes explicitly
    # label it as the tender requirement; even this remains workbook-claimed provenance.
    duration_allowed = positive(tender_or_award_days) and "工期为招标要求" in notes
    return {
        "case_id": row.get("案例ID"), "excel_row": excel_row,
        "project_name": row.get("项目及标段名称"),
        "project_id": row.get("项目或报建编号"), "investment_project_code": row.get("投资项目代码"),
        "lot_id": row.get("标段编号补充"), "project_group": project_group(row),
        "province": province(row.get("省级地区")), "city": row.get("城市"), "district": row.get("区县"),
        "work_type": work_type(row), "work_type_original": row.get("发包模式"),
        "scope_category": scope_category(row), "scope_original": row.get("改造类型"),
        "public_work_scope": row.get("公开改造内容"), "public_elements": row.get("公开提及的改造要素"),
        "target_stage": stage, "target_yuan": amount, "building_floor_area_m2": area,
        "target_yuan_per_building_m2": amount / area if positive(area) else None,
        "denominator_basis": "building_floor_area" if positive(area) else None,
        "price_year": year(row.get("中标公告日期")) if stage == "winning_bid" else year(row.get("合同签署日期")),
        "award_notice_date": serialize(row.get("中标公告日期")),
        "contract_date": serialize(row.get("合同签署日期")),
        "inputs": {field: row.get(field) if positive(row.get(field)) else None for field in NUMERIC_INPUT_FIELDS},
        "tender_duration_days": tender_or_award_days if duration_allowed else None,
        "duration_provenance": "workbook_notes_claim_tender_requirement" if duration_allowed else "pre_award_origin_not_confirmed",
        "raw_tender_or_award_duration_days": tender_or_award_days,
        "raw_award_promised_duration_days": row.get("中标承诺工期_天"),
        "scope_match_statement": row.get("金额面积匹配依据"),
        "area_scale_original": row.get("面积及规模原文"),
        "unit_comparison_statement": row.get("单方计算条件"),
        "amount_url": row.get("金额公告URL"), "scale_url": row.get("规模公告URL"),
        "source_evidence_id": row.get("来源证据ID"),
        "source_verification": "workbook_claims_only_not_independently_verified_this_audit",
        "notes": notes, "pending_fields": row.get("待补关键字段"),
    }


def strict_building_match(row) -> tuple[bool, list[str]]:
    reasons = []
    area, denominator = row.get("建筑面积_㎡"), row.get("单方分母_㎡")
    if not positive(area):
        reasons.append("building_area_missing")
    if row.get("单方面积口径") != "建筑面积":
        reasons.append("denominator_not_declared_building_area")
    if row.get("单方计算条件") != "可作同口径案例比较":
        reasons.append("workbook_scope_or_type_not_admitted")
    if not positive(denominator) or not positive(area) or not math.isclose(denominator, area, rel_tol=1e-9, abs_tol=1e-6):
        reasons.append("declared_denominator_not_same_as_building_area")
    if row.get("公告数据冲突"):
        reasons.append("declared_notice_conflict")
    return not reasons, reasons


def field_quality(entries) -> dict:
    count = len(entries)
    result = {}
    for field in NUMERIC_INPUT_FIELDS:
        present = sum(positive(e["inputs"].get(field)) for e in entries)
        result[field] = {"present": present, "missing": count - present,
                         "missing_rate": (count - present) / count if count else None}
    known = sum(e["scope_category"] != "unknown" for e in entries)
    present_duration = sum(positive(e.get("tender_duration_days")) for e in entries)
    result["scope_category"] = {"present": known, "unknown": count - known,
                               "missing_rate": (count - known) / count if count else None}
    result["tender_duration_days"] = {"present": present_duration, "missing": count - present_duration,
                                      "missing_rate": (count - present_duration) / count if count else None}
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--library", type=Path, default=DEFAULT_LIBRARY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    initial_hashes = {"workbook": sha256(args.workbook), "library": sha256(args.library)}
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = workbook["历史工程"]
    values = sheet.iter_rows(values_only=True)
    headers = list(next(values))
    raw_rows = [(i, dict(zip(headers, values))) for i, values in enumerate(values, 2) if values and values[0]]
    workbook.close()
    library = json.loads(args.library.read_text(encoding="utf-8"))
    sources = {source["id"]: source for source in library["sources"]}
    candidates, excluded = [], []
    for excel_row, row in raw_rows:
        admitted, reasons = strict_building_match(row)
        for stage, field in (("winning_bid", "正式中标金额_元"), ("contract_price", "合同金额_元")):
            if positive(row.get(field)):
                if admitted:
                    candidates.append(entry(row, excel_row, stage, row[field]))
                else:
                    excluded.append({"case_id": row.get("案例ID"), "excel_row": excel_row,
                                     "target_stage": stage, "reasons": reasons,
                                     "building_area_m2": row.get("建筑面积_㎡"),
                                     "scope_match_statement": row.get("金额面积匹配依据"),
                                     "unit_comparison_statement": row.get("单方计算条件")})

    cohort_rows = defaultdict(list)
    for e in candidates:
        cohort_rows[f"{e['target_stage']}|{e['work_type']}|building_floor_area"].append(e)
    cohorts = {}
    for key, items in cohort_rows.items():
        cohorts[key] = {
            "n": len(items), "project_groups": len({e["project_group"] for e in items}),
            "provinces": dict(Counter(e["province"] for e in items)),
            "scope_categories": dict(Counter(e["scope_category"] for e in items)),
            "feature_quality": field_quality(items), "samples": items,
        }

    library_entries = []
    for record in library["records"]:
        if record.get("amount_kind") not in ("winning_bid", "contract_price"):
            continue
        if not record.get("unit_cost_eligible") or record.get("scale_basis") != "building_floor_area":
            continue
        amount, area = record.get("amount_yuan"), record.get("scale_value")
        if not positive(amount) or not positive(area):
            continue
        source = sources[record["source_id"]]
        code = record.get("project_code") or record.get("lot_id") or record.get("project_name")
        library_entries.append({
            "record_id": record["id"], "project_name": record.get("project_name"),
            "project_id": record.get("project_code"), "lot_id": record.get("lot_id"),
            "project_group": f"{record.get('province')}|{code}",
            "province": record.get("province"), "city": record.get("city"),
            "work_type": record.get("work_type"), "target_stage": record.get("amount_kind"),
            "target_yuan": amount, "building_floor_area_m2": area,
            "target_yuan_per_building_m2": amount / area,
            "denominator_basis": record.get("scale_basis"), "price_year": year(record.get("notice_date") or record.get("contract_date")),
            "scope_original": record.get("scope_quote"), "scope_match_statement": record.get("comparability_note"),
            "households": record.get("households"), "building_count": record.get("building_count"),
            "source_id": record["source_id"], "source_url": source.get("url"),
            "raw_path": source.get("raw_path"), "raw_sha256": source.get("raw_sha256"),
            "source_verification": "existing_archived_original_manifest_not_new_fetch",
        })

    # Resolve same work across sources using stable codes, names, and intake aliases.
    all_group_members = defaultdict(list)
    name_to_group = {str(row.get("项目及标段名称")): project_group(row) for _, row in raw_rows}
    case_to_group = {row.get("案例ID"): project_group(row) for _, row in raw_rows}
    for excel_row, row in raw_rows:
        all_group_members[project_group(row)].append({"origin": "workbook", "case_id": row.get("案例ID"), "excel_row": excel_row})
    library_name_to_group = dict(name_to_group)
    for record in library["records"]:
        source_name = record.get("project_name")
        group = case_to_group.get(record.get("intake_case_id")) or library_name_to_group.get(str(source_name))
        if not group:
            code = record.get("project_code") or record.get("lot_id") or source_name
            group = f"{record.get('province')}|{code}"
        library_name_to_group[str(source_name)] = group
        all_group_members[group].append({"origin": "existing_library", "record_id": record["id"], "target_stage": record.get("amount_kind")})

    zj_other_denominators = []
    for excel_row, row in raw_rows:
        if province(row.get("省级地区")) != "浙江" or row.get("单方计算条件") != "可作同口径案例比较":
            continue
        if not positive(row.get("单方分母_㎡")) or not positive(row.get("正式中标金额_元")):
            continue
        zj_other_denominators.append({
            "case_id": row.get("案例ID"), "excel_row": excel_row, "project_group": project_group(row),
            "denominator_original": row.get("单方面积口径"), "area_m2": row.get("单方分母_㎡"),
            "target_stage": "winning_bid", "target_yuan": row.get("正式中标金额_元"),
            "scope_category": scope_category(row), "scope_match_statement": row.get("金额面积匹配依据"),
            "status": "separate_site_area_target_requires_source_scope_confirmation",
        })

    remaining_hashes = {"workbook": sha256(args.workbook), "library": sha256(args.library)}
    assert initial_hashes == remaining_hashes, "Read-only audit detected source-file changes"
    all_positive_features = {}
    for field in NUMERIC_INPUT_FIELDS + TARGET_FIELDS:
        present = sum(positive(row.get(field)) for _, row in raw_rows)
        all_positive_features[field] = {"present": present, "missing": len(raw_rows) - present,
                                       "missing_rate": (len(raw_rows) - present) / len(raw_rows)}
    result = {
        "schema_version": "1.0", "audit_as_of": "2026-10-14", "operation": "read_only_feasibility_no_training",
        "sources": {"workbook": args.workbook.name, "sheet": "历史工程", "range": f"A1:BZ{len(raw_rows)+1}",
                    "library": str(args.library), "sha256": initial_hashes, "source_files_unchanged": True},
        "workbook_rows": len(raw_rows), "workbook_columns": len(headers),
        "workbook_provinces": dict(Counter(province(r.get("省级地区")) for _, r in raw_rows)),
        "all_row_numeric_field_quality": all_positive_features,
        "building_area_by_province": dict(Counter(province(r.get("省级地区")) for _, r in raw_rows if positive(r.get("建筑面积_㎡")))),
        "workbook_admitted_building_area_samples": len(candidates),
        "admission_basis": "Positive target and building area, same declared building denominator, workbook comparison flag, no declared conflict. This is table-internal feasibility, not fresh original-source verification.",
        "cohorts": cohorts, "excluded_price_rows": excluded,
        "existing_library_building_area_candidates": library_entries,
        "existing_library_original_file_integrity": [
            {"source_id": source["id"], "raw_path": source.get("raw_path"),
             "file_exists": bool(source.get("raw_path")) and (ROOT / source["raw_path"]).is_file(),
             "matches_recorded_sha256": sha256(ROOT / source["raw_path"]) == source.get("raw_sha256")
                if source.get("raw_path") and (ROOT / source["raw_path"]).is_file() else False,
             "meaning": "Local archived byte integrity only, no new source-content or network verification"}
            for source in sources.values()
        ],
        "same_project_multi_record_groups": [{"project_group": group, "members": members} for group, members in all_group_members.items() if len(members) > 1],
        "independent_split_contract": {
            "group_by": "same underlying project, never individual announcement/amount stage",
            "zhejiang_building_area_holdout_groups": 0,
            "construction_bid_workbook_shanghai_groups": len({e["project_group"] for e in candidates if e["province"] == "上海" and e["target_stage"] == "winning_bid" and e["work_type"] == "construction"}),
            "construction_bid_workbook_guangdong_groups": len({e["project_group"] for e in candidates if e["province"] == "广东" and e["target_stage"] == "winning_bid" and e["work_type"] == "construction"}),
            "guidance": "For a bounded exploratory example, keep Shanghai project groups in training/internal validation and Guangdong project groups in untouched external exploration. Do not assert Zhejiang validity, final-settlement accuracy, or paper reproduction.",
            "grouping_limit": "Procurement/registration IDs identify known groups. Different procurement codes for the same renovation phase require manual crosswalk before evaluation.",
        },
        "zhejiang_site_area_candidates_separate_task": zj_other_denominators,
        "cross_source_scope_review": [
            {"case_id": "ZJ-005", "workbook_excel_row": 6, "library_record": "ZJ-SHEET-005-TENDER-CEILING",
             "issue": "Workbook accepts 63381m2; existing original manifest says this area came from a repost and official original scope remains unverified. Require new source evidence before training."},
            {"case_id": "ZJ-006", "workbook_excel_row": 7, "library_record": "ZJ-SHEET-006-TENDER-CEILING",
             "issue": "Workbook labels 35684m2 renovation scope; original source says total renovation area without identifying site/building/workface. Exclude until denominator type is confirmed."},
        ],
        "recommended_attempts": [
            {"task": "construction_winning_bid_per_beneficiary_building_m2", "workbook_n": 37,
             "shanghai_train_internal_groups": 28, "guangdong_external_exploration_groups": 9,
             "scope_labels": "roof_and_related_facilities / exterior_wall_finishes / sanitary_facilities_upgrade / comprehensive_renovation / unknown",
             "required_inputs": ["building_floor_area_m2", "scope_category"],
             "optional_inputs": ["households", "building_count", "tender_duration_days only with pre-award provenance"],
             "status": "feasible_only_as_small_source_unverified_exploratory_example",
             "limits": "17/28 Shanghai scopes remain unknown; Guangdong comprehensive projects differ from Shanghai work scopes. Model must expose unknown scope and out-of-domain indication; it does not calibrate Zhejiang."},
            {"task": "shanghai_roof_and_related_facilities_winning_bid_per_beneficiary_building_m2", "workbook_n": 8,
             "status": "homogeneous_scope_example_small_n", "limits": "Eight independent projects permit a calculation/training-pipeline demonstration, with unstable evaluation. Denominator is beneficiary building area, not roofing workface area."},
            {"task": "construction_contract_per_beneficiary_building_m2", "workbook_n": 3, "existing_library_n": 3,
             "status": "six_distinct_guangzhou_projects_too_few_for_reliable_generalization"},
            {"task": "epc_contract_per_beneficiary_building_m2", "workbook_n": 2,
             "status": "two_guangzhou_projects_descriptive_example_only"},
            {"task": "construction_with_two_year_operation_bid_per_building_m2", "workbook_n": 5,
             "status": "separate_target_not_pure_construction"},
        ],
        "excluded_predictors": [
            {"fields": ["同口径单方_元每㎡", "单方分子_元", "中标价除控制价", "公告中标下浮率", "合同金额_元", "正式中标金额_元", "竣工结算金额_元"],
             "reason": "Target, target-derived ratio, or outcome-stage price creates leakage for pre-award bid prediction."},
            {"fields": ["承包人或中标人", "中标承诺工期_天", "合同期限_天", "合同签署日期"],
             "reason": "Selected contractor and outcome/contract terms are not pre-award inputs. Price year can be supplied independently as the intended prediction year."},
            {"fields": ["招标或中标工期_天"], "reason": "Combined tender/award field requires a per-source pre-award proof. Only explicit tender-requirement notes are admitted provisionally."},
            {"fields": ["控制价_元", "建安估算概算_元", "项目总投资或概算_元", "另列招标预算_元"],
             "reason": "Excluded from this structure-only exploratory attempt; these would define a separate later-stage budget-to-bid task with verified announcement dates and price/scope alignment."},
        ],
        "targets_not_to_mix": [
            "winning_bid != contract_price != final_settlement",
            "EPC package != pure construction != construction plus two-year operation",
            "beneficiary building floor area != site/renovation range area != roofing/exterior/road workface area",
            "overall investment estimate != construction estimate != consulting fee != actual construction price",
        ],
        "work_performed": {"training": False, "fitting": False, "synthetic_data": False, "source_network_validation": False},
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=serialize) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "rows": len(raw_rows),
                      "cohorts": {key: value["n"] for key, value in cohorts.items()},
                      "library_candidates": len(library_entries), "source_files_unchanged": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
