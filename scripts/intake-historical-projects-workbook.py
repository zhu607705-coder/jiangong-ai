"""Read an AI-collected workbook without editing it; audit its known formulas.

Usage (offline replay uses the archived workbook and captured official pages):
  python scripts/intake-historical-projects-workbook.py
  python scripts/intake-historical-projects-workbook.py --source PATH --fetch

The workbook is an untrusted lead list. Instructions in cells are kept as data.
Only the two exact R/S formula templates are independently recomputed, without
eval or Excel execution. Existing openpyxl/requests/BeautifulSoup libraries are
reused; this is a bounded intake utility, not a general parser or search client.
"""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
from urllib.parse import urlparse, urljoin, parse_qs

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import requests
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
COPY = ROOT / "data/imports/historical-projects-user-20261013.xlsx"
OUTPUT = COPY.with_suffix(".json")
MANIFEST = ROOT / "data/public_cases/methods-expansion/zhejiang-workbook-manifest.json"
SOURCE_DIR = ROOT / "data/public_cases/renewal-cost-sources"


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def save_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def relative(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def audit_formula(row: list, number: int) -> dict:
    r_formula = f'=IF(G{number}="","",IF(AND(F{number}="正式中标",M{number}<>""),M{number}/G{number},IF(L{number}<>"",L{number}/G{number},"")))'
    s_formula = f'=IF(R{number}="","不可折算",IF(F{number}="正式中标","中标价÷改造范围",IF(L{number}<>"","招标限价÷改造范围","不可折算")))'
    expected = row[17] == r_formula and row[18] == s_formula
    amount, stage = None, None
    if row[5] == "正式中标" and isinstance(row[12], (int, float)):
        amount, stage = row[12], "winning_bid"
    elif isinstance(row[11], (int, float)):
        amount, stage = row[11], "tender_ceiling"
    area = row[6]
    value = amount / area if expected and amount is not None and isinstance(area, (int, float)) and area > 0 else None
    label = "不可折算" if value is None else ("中标价÷改造范围" if row[5] == "正式中标" else "招标限价÷改造范围")
    return {"known_templates_matched": expected, "independent_R": value, "independent_S": label,
            "numerator_stage_from_workbook": stage, "denominator_field": "G:改造范围面积㎡",
            "meaning": "综合改造范围金额强度；不是道路、管道、立面等工程分项单价",
            "verified_observation": False}


def archive_source(item: tuple[str, str], fetch: bool) -> dict:
    source_id, url = item
    meta_path = SOURCE_DIR / f"{source_id}.json"
    if not fetch and meta_path.exists():
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        for field, hash_field in (("raw_path", "raw_sha256"), ("text_path", "text_sha256")):
            if digest((ROOT / meta[field]).read_bytes()) != meta[hash_field]:
                raise ValueError(f"Source hash changed: {meta[field]}")
        raw_soup = BeautifulSoup((ROOT / meta["raw_path"]).read_bytes().decode("utf-8"), "html.parser")
        article_title = raw_soup.find("meta", attrs={"name": "ArticleTite"})
        if article_title and article_title.get("content"):
            meta["title"] = article_title["content"]
            save_json(meta_path, meta)
        return meta
    if urlparse(url).hostname != "dzjy.lssggzy.lishui.gov.cn":
        raise ValueError("Only the reviewed official Lishui source domain can be captured")
    response = requests.get(url, timeout=30)
    response.raise_for_status()
    if urlparse(response.url).hostname != "dzjy.lssggzy.lishui.gov.cn":
        raise ValueError("Unexpected source redirect")
    soup = BeautifulSoup(response.content.decode("utf-8"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = " ".join(soup.get_text(" ", strip=True).split())
    lot_match = re.search(r"项目编号[：:]\s*([A-Z0-9]+)", text)
    date_match = re.search(r"信息时间[：:]\s*(\d{4}-\d{2}-\d{2})", text)
    if not lot_match or not date_match:
        raise ValueError(f"Official notice structure missing: {url}")
    raw_path, text_path = SOURCE_DIR / f"{source_id}.html", SOURCE_DIR / f"{source_id}.txt"
    SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    raw_path.write_bytes(response.content)
    text_path.write_text(text + "\n", encoding="utf-8")
    article_title = soup.find("meta", attrs={"name": "ArticleTite"})
    title = article_title.get("content") if article_title else (soup.title.get_text(strip=True) if soup.title else source_id)
    meta = {"id": source_id, "title": title,
            "requested_url": url, "final_url": response.url, "url": response.url,
            "province": "浙江", "published_at": date_match[1],
            "fetched_at": datetime.now(timezone.utc).isoformat(), "http_status": response.status_code,
            "raw_format": "html", "content_type": response.headers.get("Content-Type"),
            "raw_path": relative(raw_path), "text_path": relative(text_path),
            "raw_sha256": digest(response.content), "text_sha256": digest(text_path.read_bytes()),
            "source_status": "fetched_official_public_source",
            "extraction": "BeautifulSoup script/style removed; whitespace normalized", "lot_id": lot_match[1]}
    save_json(meta_path, meta)
    return meta


def quote(text: str, pattern: str) -> str:
    match = re.search(pattern, text)
    if not match:
        raise ValueError(f"Reviewed quote not found: {pattern}")
    return match[0]


def archive_tender_document(notice: dict, fetch: bool) -> dict:
    """Follow the public attachment form found in the official notice, then hash PDF bytes."""
    source_id = notice["id"] + "-document"
    meta_path = SOURCE_DIR / f"{source_id}.json"
    if not fetch and meta_path.exists():
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        for key, hash_key in (("raw_path", "raw_sha256"), ("text_path", "text_sha256")):
            if digest((ROOT / meta[key]).read_bytes()) != meta[hash_key]:
                raise ValueError(f"Source hash changed: {meta[key]}")
        return meta
    soup = BeautifulSoup((ROOT / notice["raw_path"]).read_bytes().decode("utf-8"), "html.parser")
    anchor = next(tag for tag in soup.find_all("a") if "招标文件正文" in tag.get_text())
    form_url = urljoin(notice["url"], anchor["href"])
    form = requests.get(form_url, timeout=30)
    form.raise_for_status()
    # The site's public download form declares this action with these URL parameters.
    if 'ztbAttachDownloadAction.action?cmd=getContent&attachGuid=' not in form.text:
        raise ValueError("Official download form changed; review before following it")
    params = parse_qs(urlparse(form_url).query)
    endpoint = urljoin(form_url, "ztbAttachDownloadAction.action")
    if urlparse(endpoint).hostname != "dzjy.lssggzy.lishui.gov.cn":
        raise ValueError("Unexpected attachment host")
    response = requests.post(endpoint, params={"cmd": "getContent", **{key: params[key][0] for key in ("attachGuid", "appUrlFlag", "siteGuid")}}, timeout=30)
    response.raise_for_status()
    if not response.content.startswith(b"%PDF-"):
        raise ValueError("Official attachment did not return a PDF")
    raw_path = SOURCE_DIR / f"{source_id}.pdf"
    text_path = SOURCE_DIR / f"{source_id}.txt"
    raw_path.write_bytes(response.content)
    reader = PdfReader(raw_path)
    text = "\n".join(f"[PDF page {i + 1}]\n" + " ".join((page.extract_text() or "").split()) for i, page in enumerate(reader.pages)) + "\n"
    text_path.write_text(text, encoding="utf-8")
    meta = {"id": source_id, "title": notice["title"].removesuffix("招标公告") + "施工招标文件",
            "requested_url": form_url, "final_url": response.url, "url": form_url,
            "landing_notice_url": notice["url"], "linked_notice_source_id": notice["id"],
            "province": "浙江", "published_at": notice["published_at"],
            "fetched_at": datetime.now(timezone.utc).isoformat(), "http_status": response.status_code,
            "raw_format": "pdf", "content_type": response.headers.get("Content-Type"),
            "raw_path": relative(raw_path), "text_path": relative(text_path),
            "raw_sha256": digest(response.content), "text_sha256": digest(text_path.read_bytes()),
            "source_status": "fetched_official_public_source", "retrieval_method": "GET official public attachment form then POST its declared download action",
            "extraction": "pypdf page-wise text with whitespace normalization and explicit one-based PDF page labels",
            "pages": len(reader.pages), "lot_id": None, "linked_notice_lot_id": notice["lot_id"],
            "reviewed_pages": [3, 9]}
    save_json(meta_path, meta)
    return meta


def build_records(sources: list[dict]) -> list[dict]:
    by_id = {source["id"]: source for source in sources}
    records = []

    def add(source_id: str, case: str, kind: str, amount: int, amount_pattern: str,
            scope_pattern: str, note: str, *, scale=None, scale_pattern=None,
            scale_basis="unknown", eligible=False, bidder=None, provisional=None) -> None:
        source = by_id[source_id]
        text = (ROOT / source["text_path"]).read_text(encoding="utf-8")
        amount_quote = quote(text, amount_pattern)
        scope_quote = quote(text, scope_pattern)
        amount_number = re.search(r"(?:[0-9]+(?:\.[0-9]+)?)\s*(?:万元|元)", amount_quote)
        if amount_number:
            numeric = float(re.search(r"[0-9]+(?:\.[0-9]+)?", amount_number[0])[0])
            expected = numeric * (10000 if "万元" in amount_number[0] else 1)
            if not math.isclose(expected, amount, abs_tol=0.01):
                raise ValueError(f"Amount differs from its quote: {source_id}")
        project_match = re.search(r"项目代码\s*[：:]\s*([0-9-]+)", text)
        record = {"id": f"ZJ-SHEET-{case}-{kind.upper().replace('_', '-')}",
                  "source_id": source_id, "project_name": scope_quote, "project_code": project_match[1] if project_match else None,
                  "lot_id": source["lot_id"], "province": "浙江", "city": "丽水",
                  "district": "遂昌县" if "遂昌县" in text else "莲都区", "work_type": "construction",
                  "amount_kind": kind, "amount_yuan": amount, "amount_original": amount_number[0] if amount_number else f"{amount}元",
                  "amount_quote": amount_quote, "scope_quote": scope_quote,
                  "scale_value": scale, "scale_unit": "m2" if scale is not None else None,
                  "scale_quote": quote(text, scale_pattern) if scale_pattern else None, "scale_basis": scale_basis,
                  "notice_date": source["published_at"], "contract_date": None,
                  "actual_settled_cost": False, "unit_cost_eligible": eligible,
                  "amount_basis": {"tender_ceiling": "construction_tender_control", "winning_bid": "construction_award", "candidate_bid": "candidate_bid_not_award"}[kind],
                  "winning_bidder": bidder, "scope_level": "construction_lot", "comparability_note": note,
                  "intake_case_id": f"ZJ-{case}", "lead_source_class": "user_ai_collected",
                  "source_class": "official_public_verified", "provisional_sum_yuan": provisional}
        if kind == "candidate_bid":
            record["candidate_count_same_quote"] = 5
        records.append(record)

    add("zj-sheet-001-tender-document", "001", "tender_ceiling", 18997007,
        r"最高投标限价：18997007元（其中暂列金额（含税）872000元，[^。]+?下浮）", r"遂昌县妙高街道城镇老旧小区改造项目（一期）——腾龙一期区块",
        "招标正文第9页明确最高投标限价，含不下浮暂列金872000元；第3页面积为占地改造范围，只作综合范围金额强度。",
        scale=92815, scale_pattern=r"本次\s*改造范围占地约92815平方米[^。]+", scale_basis="renovation_site_area", eligible=True, provisional=872000)
    add("zj-sheet-001-award", "001", "winning_bid", 17333131,
        r"中标价： 人民币 [^）]+17333131\.00元）", r"遂昌县妙高街道城镇老旧小区改造项目（一期）——腾龙一期区块",
        "正式中标通知书。此原件未给面积；与同标段招标公告的面积联动另存 intake.source_link_pairings，不能伪装成单页原文。", bidder="浙江立骏建设有限公司")
    add("zj-sheet-002-award", "002", "winning_bid", 15012681,
        r"中标价 : 大写 ： [^）]+15012681\.00元 ）", r"丽水市莲都区东升北区1-3栋、南区13栋及周边老旧小区改造工程",
        "正式中标通知书，载明中标下浮率11.9%；未给可作分母的改造面积。", bidder="浙江华歌建设有限公司")
    add("zj-sheet-003-tender", "003", "tender_ceiling", 8314400,
        r"本次暂定的最高限价（工程费用）为\s*831\.44\s*万元[^）]+）", r"丽水市莲都区东升南区及周边老旧小区改造工程",
        "暂定最高限价，最终以审核施工图预算为准；13477.831㎡只是道路修复面积，金额还含管网等，禁止相除为道路单价。",
        scale=13477.831, scale_pattern=r"道路改造\s*[：:]\s*修复面积\s*13477\.831\s*平方米", scale_basis="road_area")
    add("zj-sheet-004-tender-document", "004", "tender_ceiling", 7732717,
        r"最高投标限价 7732717元（[^）]+）", r"遂昌县妙高街道城镇老旧小区改造项目（一期）——叶坦区块",
        "招标正文第9页明确勾选最高投标限价，含不下浮暂列金359700元；第3页53021㎡为项目范围，非建筑面积。",
        scale=53021, scale_pattern=r"项目范围约53021平方米[^：]+：", scale_basis="renovation_site_area", eligible=True, provisional=359700)
    add("zj-sheet-005-amendment", "005", "tender_ceiling", 9667936,
        r"本次招标建安工程造价： 9667936元【.+?最高投标限价下浮： 6\.0%-12\.0% （[^）]+）", r"遂昌县妙高街道城镇老旧小区改造项目（一期）——渡船头区块工程",
        "补充修改通知的招标控制价，含不下浮暂列金327000元；原表63381㎡来自转载，官方原始面积尚待核实，面积不入核实记录。", provisional=327000)
    add("zj-sheet-006-tender", "006", "tender_ceiling", 11005241,
        r"最高投标限价\s*[：:]\s*11005241\s*元", r"丽水市莲都区金山新村及周边老旧小区改造工程",
        "原文写总改造面积35684㎡、18幢，未明示占地或建筑面积；面积保留但scale_basis=unknown，先排除同口径金额强度统计。",
        scale=35684, scale_pattern=r"项目总改造面积\s*35684\s*平方米[^。]+", scale_basis="unknown")
    add("zj-sheet-007-candidate", "007", "candidate_bid", 14972288,
        r"预中标价 : 大写 [^）]+14972288 元）", r"丽水市莲都区正达阳光城老旧小区改造工程",
        "5名不分排名中标候选人均报14972288元；保存为一条标段候选报价，不算5个独立工程、不视作正式中标或结算。")
    return records


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=COPY)
    parser.add_argument("--fetch", action="store_true", help="Fetch reviewed official notices; otherwise replay archived pages")
    args = parser.parse_args()
    source_bytes = args.source.read_bytes()
    sha = digest(source_bytes)
    COPY.parent.mkdir(parents=True, exist_ok=True)
    if args.source.resolve() != COPY.resolve():
        if COPY.exists() and digest(COPY.read_bytes()) != sha:
            raise ValueError("Archived original already exists with a different SHA256")
        shutil.copyfile(args.source, COPY)
    before_hash = digest(args.source.read_bytes())
    formulas = load_workbook(COPY, read_only=True, data_only=False)
    cached = load_workbook(COPY, read_only=True, data_only=True)
    sheets = []
    for sheet in formulas:
        cached_sheet = cached[sheet.title]
        rows = []
        for cells in sheet.iter_rows():
            rows.append({"source_row": cells[0].row, "cells": [
                {"coordinate": cell.coordinate, "value": None if cell.data_type == "f" else cell.value,
                 "formula": cell.value if cell.data_type == "f" else None,
                 "cached_value": cached_sheet.cell(cell.row, cell.column).value,
                 "data_type": cell.data_type} for cell in cells]})
        sheets.append({"name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "rows": rows})
    main_sheet = formulas["官方工程案例"]
    lead_rows, items = [], []
    source_keys = {(2, 14): "001-tender", (2, 15): "001-award", (3, 14): "002-tender", (3, 15): "002-award",
                   (4, 14): "003-tender", (5, 14): "004-tender", (6, 14): "005-amendment", (7, 14): "006-tender", (8, 14): "007-candidate"}
    headers = [cell.value for cell in next(main_sheet.iter_rows())]
    for cells in main_sheet.iter_rows(min_row=2):
        row = [cell.value for cell in cells]
        number = cells[0].row
        audit = audit_formula(row, number)
        cached_r = cached[main_sheet.title].cell(number, 18).value
        cached_s = cached[main_sheet.title].cell(number, 19).value
        audit["R_cached_matches"] = (cached_r is None and audit["independent_R"] is None) or (
            isinstance(cached_r, (int, float)) and audit["independent_R"] is not None and math.isclose(cached_r, audit["independent_R"], abs_tol=1e-8))
        audit["S_cached_matches"] = cached_s == audit["independent_S"]
        lead_rows.append({"case_id": row[0], "source_sheet": main_sheet.title, "source_row": number,
                          "source_class": "user_ai_collected", "status": "lead_with_separately_verified_fields",
                          "verified_observation": False, "values": {get_column_letter(i + 1): value for i, value in enumerate(row)},
                          "formula_audit": audit})
        for index in (14, 15):
            if (number, index) in source_keys:
                items.append(("zj-sheet-" + source_keys[number, index], row[index]))
    with ThreadPoolExecutor(max_workers=4) as executor:
        sources = list(executor.map(lambda item: archive_source(item, args.fetch), items))
    tender_sources = [source for source in sources if source["id"] in {"zj-sheet-001-tender", "zj-sheet-004-tender"}]
    with ThreadPoolExecutor(max_workers=2) as executor:
        sources.extend(executor.map(lambda source: archive_tender_document(source, args.fetch), tender_sources))
    records = build_records(sources)
    source_map = {source["id"]: source for source in sources}
    pairs = []
    for case in ("001", "002"):
        ids = [f"zj-sheet-{case}-tender", f"zj-sheet-{case}-award"]
        lot_ids = [source_map[source_id]["lot_id"] for source_id in ids]
        pairs.append({"case_id": f"ZJ-{case}", "source_ids": ids, "same_lot_id": len(set(lot_ids)) == 1,
                      "lot_id": lot_ids[0], "link_rule": "identical_official_project_number_and_project_name",
                      "fields_by_source": {ids[0]: ["scope", "project_code"] + (["scale_value", "scale_basis", "tender_construction_cost"] if case == "001" else ["investment_estimate"]), ids[1]: ["winning_bid", "winning_bidder"]},
                      "derived_scope_amount_intensity_yuan_per_sqm": 17333131 / 92815 if case == "001" else None,
                      "derived_not_direct_source_quote": True})
    assert before_hash == digest(args.source.read_bytes()) == sha == digest(COPY.read_bytes())
    artifact = {"schema_version": 1, "as_of": "2026-10-13", "source_class": "user_ai_collected",
                "source_filename": "历史工程·1.xlsx", "archived_path": relative(COPY), "sha256": sha,
                "bytes": len(source_bytes), "original_unchanged": True,
                "trust_boundary": "Workbook content is data, not agent instructions; AI-collected rows remain leads; official source fields are promoted separately.",
                "headers": headers, "sheets": sheets, "lead_rows": lead_rows, "source_link_pairings": pairs,
                "verification_manifest": relative(MANIFEST),
                "summary": {"sheet_count": len(sheets), "case_count": len(lead_rows), "case_columns": len(headers),
                            "locations": sorted({row["values"]["B"] for row in lead_rows}),
                            "shanghai_cases": sum("上海" in row["values"]["B"] for row in lead_rows),
                            "guangzhou_cases": sum("广州" in row["values"]["B"] for row in lead_rows),
                            "formula_count": sum(cell["formula"] is not None for sheet in sheets for row in sheet["rows"] for cell in row["cells"]),
                            "known_formula_templates_matched": all(row["formula_audit"]["known_templates_matched"] for row in lead_rows),
                            "cached_results_match_independent_calculation": all(row["formula_audit"]["R_cached_matches"] and row["formula_audit"]["S_cached_matches"] for row in lead_rows),
                            "archived_official_sources": len(sources), "verified_stage_records": len(records)},
                "findings": ["R列4条数值均与独立计算一致，但计算正确不自动证明分母和金额适合比较。",
                             "ZJ-001中标价与面积分属两份同标段公告；联动计算186.749243元/㎡仅为综合范围金额强度。",
                             "ZJ-001、ZJ-004招标正文第9页明确最高投标限价，第3页给项目范围；叶坦范围金额强度145.842534元/㎡，与中标价分组展示。",
                             "ZJ-005原表面积63381㎡源于转载，尚未升级为官方核实面积。",
                             "ZJ-006原文总改造面积未明确占地/建筑，scale_basis=unknown。",
                             "ZJ-007五名候选人同报14972288元，按一标段候选报价保存，不当正式中标或独立样本。"],
                "reproduce": "python scripts/intake-historical-projects-workbook.py"}
    save_json(MANIFEST, {"as_of": "2026-10-13", "sources": sources, "records": records,
                         "intake": relative(OUTPUT), "verification": "Exact source quotes and raw/text SHA256 checked by intake script; no construction settlement records in this batch."})
    save_json(OUTPUT, artifact)
    formulas.close()
    cached.close()
    print(json.dumps(artifact["summary"], ensure_ascii=False))


if __name__ == "__main__":
    main()
