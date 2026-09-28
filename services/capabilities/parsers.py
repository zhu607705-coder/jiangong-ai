"""The binary-format branch of the one ParserPort; all processing is local."""
from pathlib import Path
from zipfile import ZipFile
from hashlib import sha256
import os

import pandas as pd
from docx import Document
from vendor.cre.parsers import extract_pages, prepare_table, build_table_preview


def parse_document(file_path: str, file_name: str, document_id: str) -> dict:
    root = (Path(os.environ.get("DATA_DIR", "data/local")) / "blobs").resolve()
    path = Path(file_path).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError("Parser input is outside the authorized blob directory")
    if path.stat().st_size > 10 * 1024 * 1024:
        raise ValueError("File exceeds 10 MB parser limit")
    ext = Path(file_name).suffix.lower()
    if ext not in {".pdf", ".docx", ".xlsx"}:
        raise ValueError("Unsupported binary format")
    if ext in {".docx", ".xlsx"}:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                raise ValueError("Archive expands beyond the parser budget")
    segments, tables, notes = [], [], []
    if ext == ".pdf":
        pages = extract_pages(str(path))
        segments = [{"locator": f"page:{i + 1}", "text": text} for i, text in enumerate(pages) if text.strip()]
        notes.append("最多解析前 50 页，页码定位保留；扫描页需要 OCR，不能被解释为没有内容。")
        status = "extracted" if segments else "ocr_required"
        if len(segments) < len(pages): notes.append("部分页面没有可提取文本，可能需要 OCR。")
    elif ext == ".docx":
        doc = Document(path)
        segments = [{"locator": f"paragraph:{i + 1}", "text": paragraph.text} for i, paragraph in enumerate(doc.paragraphs) if paragraph.text.strip()]
        for index, table in enumerate(doc.tables[:50]):
            rows = [[cell.text for cell in row.cells] for row in table.rows[:251]]
            if not rows: continue
            segments.append({"locator": f"table:{index + 1}", "text": "\n".join(" | ".join(row) for row in rows)})
            headers = [f"{value or 'Column'} ({i + 1})" for i, value in enumerate(rows[0])]
            tables.append({"tableId": f"{document_id}-T{index + 1}", "label": f"Word table {index + 1}", "rowCount": len(rows) - 1, "columnCount": len(headers), "columns": headers, "rows": [{"rowNumber": i + 2, "values": dict(zip(headers, row))} for i, row in enumerate(rows[1:])]})
        status = "extracted" if segments else "empty"
    else:
        workbook = pd.ExcelFile(path, engine="openpyxl")
        if len(workbook.sheet_names) > 30: raise ValueError("Workbook exceeds 30-sheet budget")
        for sheet in workbook.sheet_names:
            raw = pd.read_excel(workbook, sheet_name=sheet, header=None, nrows=5001)
            if raw.empty: continue
            table = build_table_preview(prepare_table(raw, sheet))
            table["tableId"] = f"{document_id}-{len(tables) + 1}"
            tables.append(table)
            for row in table["rows"]:
                segments.append({"locator": f"sheet:{sheet}/row:{row['rowNumber']}", "text": " | ".join(f"{key}={value}" for key, value in row["values"].items())})
        status = "extracted" if segments else "empty"
        notes.append("公式只读取文件缓存值，不执行宏或外部链接；最多 30 表、每表 5000 行，展示前 250 行。")
    text = "\n\n".join(f"[{s['locator']}] {s['text']}" for s in segments)
    if len(text) > 100000: notes.append("文本超过 100000 字符，研判输入已截断，原文件仍保留。")
    return {"status": status, "parser": "CRE generic PDF/Excel + python-docx", "source_hash": sha256(path.read_bytes()).hexdigest(),
            "text": text[:100000], "segments": segments[:1000], "profile": {"tableCount": len(tables), "tables": tables},
            "notes": notes + ["解析内容仍是待核证据，不能自动成为法规、组织约束或工程结论。"]}
